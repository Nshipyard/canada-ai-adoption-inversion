#!/usr/bin/env python3
"""Build the normalized datasets for canada-ai-adoption-inversion.

US side: extracted from the U.S. Census Bureau's own "AI Core Questions" file
(legacy wording series, "in producing goods or services"), archived copy at
scripts/sources/btos_ai_core.xlsx. Live file: https://www.census.gov/hfp/btos/downloads/
("AI Core Questions.xlsx"; the archived 2026-08-11 copy still carries the original
wording series). Question change documented in scripts/sources/btos_ai_wording.pdf
(Census, Dec 3 2025): wording changed Nov 17, 2025, new series from Dec 4, 2025.

Canada side: figures transcribed from the Statistics Canada analytical publication
"Analysis on artificial intelligence use by businesses in Canada, second quarter of 2025"
(11-621-m2025008, released 2025-06-16) and the Q2 2026 edition (11-621-m2026010).
The publication cites CSBC Table 33-10-1004-01. The StatCan WDS full-table download
and the /n1/tbl/csv static zip both served stale/wrong content for this cube at build
time (2026-10-09), so sector and headline figures are transcribed from the publication
text, which is the primary source. Firm-size cuts (Table 33-10-1005) were not
retrievable; the site shows the US firm-size panel only, clearly labeled.

Outputs: public/data/*.json and *.csv (MIT licensed).
"""
import csv
import json
import re
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "scripts" / "sources" / "btos_ai_core.xlsx"
OUT = ROOT / "public" / "data"
OUT.mkdir(parents=True, exist_ok=True)

PULL_DATE = "2026-10-09"
CENSUS_FILE_URL = "https://www.census.gov/hfp/btos/downloads/"
CENSUS_WORDING_PDF = "https://www.census.gov/hfp/btos/downloads/AI%20Question%20Wording%20Updates.pdf"
STATCAN_Q2_2025 = "https://www150.statcan.gc.ca/n1/pub/11-621-m/11-621-m2025008-eng.htm"
STATCAN_Q2_2026 = "https://www150.statcan.gc.ca/n1/pub/11-621-m/11-621-m2026010-eng.htm"


def pct(s):
    if s is None:
        return None
    s = str(s).strip()
    if s in ("", "S", "None"):
        return None
    m = re.match(r"([\d.]+)%", s)
    return float(m.group(1)) if m else None


wb = openpyxl.load_workbook(SRC, read_only=True, data_only=True)

# --- National legacy series (Question 7 = current use, Answer 1 = Yes) ---
ws = wb["National Estimates"]
hdr = [str(c) for c in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
codes = [c for c in hdr[4:] if c and c != "None"]
wsd = wb["Collection and Reference Dates"]
dmap = {}
for r in list(wsd.iter_rows(min_row=2, values_only=True)):
    if r[0] and str(r[0]).isdigit():
        dmap[str(r[0])] = {"col_start": str(r[1]), "ref_start": str(r[3]), "ref_end": str(r[4])}
series = []
for row in ws.iter_rows(min_row=2, values_only=True):
    r = list(row)
    if str(r[0]) == "7" and str(r[2]) == "1":
        vals = list(r[4:])
        for c, v in zip(codes, vals):
            p = pct(v)
            if p is not None and c in dmap:
                series.append({"fortnight": c, **dmap[c], "pct": p})
        break
series.sort(key=lambda x: x["fortnight"])
assert len(series) > 50, f"expected 50+ fortnights, got {len(series)}"

# September reference points: the fortnight whose reference period is fully inside
# September (the labeled comparison; avoids fortnights straddling month boundaries)
sept = {}
for pt in series:
    ms = re.match(r"(\d{2})/\d{2}/(\d{4})", pt["ref_start"])
    me = re.match(r"(\d{2})/\d{2}/(\d{4})", pt["ref_end"])
    if ms and me and ms.group(1) == "09" and me.group(1) == "09" and ms.group(2) == me.group(2):
        sept[ms.group(2)] = pt["pct"]
        sept[ms.group(2) + "_f"] = pt["fortnight"]
for yr, want in (("2023", 3.7), ("2024", 6.0), ("2025", 10.0)):
    got = sept.get(yr)
    assert got is not None and abs(got - want) < 0.05, f"Sept {yr}: expected {want}, got {got}"
print("US legacy series verified: Sept 2023", sept["2023"], "/ Sept 2024", sept["2024"], "/ Sept 2025", sept["2025"])

# --- Sector cuts, fortnight 202520 (ref 09/21/2025) ---
SECTOR_NAMES = {
    "11": "Agriculture, forestry, fishing and hunting", "21": "Mining, quarrying, and oil and gas extraction",
    "22": "Utilities", "23": "Construction", "31": "Manufacturing", "42": "Wholesale trade",
    "44": "Retail trade", "48": "Transportation and warehousing", "51": "Information",
    "52": "Finance and insurance", "53": "Real estate and rental and leasing",
    "54": "Professional, scientific and technical services", "55": "Management of companies and enterprises",
    "56": "Administrative and support, waste management and remediation services",
    "61": "Educational services", "62": "Health care and social assistance",
    "71": "Arts, entertainment and recreation", "72": "Accommodation and food services",
    "81": "Other services (except public administration)", "XX": "Multi-unit companies",
}
ws = wb["Sector Estimates"]
hdr = [str(c) for c in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
idx = hdr.index("202520")
sectors = []
for row in ws.iter_rows(min_row=2, values_only=True):
    r = list(row)
    code = str(r[0])
    if code in SECTOR_NAMES and str(r[1]) == "7" and str(r[3]) == "1":
        sectors.append({"naics": code, "sector": SECTOR_NAMES[code], "pct": pct(r[idx])})
assert len(sectors) == 20

# --- Firm-size cuts, fortnight 202520 ---
BANDS = {"A": "1-4", "B": "5-9", "C": "10-19", "D": "20-49", "E": "50-99", "F": "100-249", "G": "250+"}
ws = wb["Employment Size Estimates"]
hdr = [str(c) for c in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
idx = hdr.index("202520")
firmsize = []
for row in ws.iter_rows(min_row=2, values_only=True):
    r = list(row)
    if str(r[0]) in BANDS and str(r[1]) == "7" and str(r[3]) == "1":
        firmsize.append({"band": BANDS[str(r[0])], "pct": pct(r[idx])})
order = ["1-4", "5-9", "10-19", "20-49", "50-99", "100-249", "250+"]
firmsize.sort(key=lambda x: order.index(x["band"]))
assert len(firmsize) == 7

# --- StatCan side (transcribed from the primary publications, see docstring) ---
canada = {
    "question": "Used artificial intelligence to produce goods or deliver services over the 12 months preceding the survey",
    "survey": "Canadian Survey on Business Conditions",
    "source_publication": STATCAN_Q2_2025,
    "source_table_cited": "33-10-1004-01",
    "sample": {"invited": 21357, "respondents": 9103},
    "headline": [
        {"quarter": "Q2 2024", "collection": "2024-04-02 to 2024-05-06", "pct": 6.1},
        {"quarter": "Q2 2025", "collection": "2025-04-01 to 2025-05-05", "pct": 12.2},
        {"quarter": "Q2 2026", "collection": "2026-04-01 to 2026-05-06", "pct": 19.2,
         "source": STATCAN_Q2_2026, "note": "latest available at build time"},
    ],
    "sectors_q2_2025": [
        {"sector": "Information and cultural industries", "pct_2025": 35.6, "pct_2024": 20.9},
        {"sector": "Professional, scientific and technical services", "pct_2025": 31.7, "pct_2024": 13.7},
        {"sector": "Finance and insurance", "pct_2025": 30.6, "pct_2024": 10.9},
        {"sector": "Accommodation and food services", "pct_2025": 1.5, "pct_2024": None},
        {"sector": "Agriculture, forestry, fishing and hunting", "pct_2025": 1.8, "pct_2024": None},
        {"sector": "Transportation and warehousing", "pct_2025": 1.8, "pct_2024": None},
    ],
    "firm_size_note": "Firm-size cuts are published in CSBC Table 33-10-1005, which was not retrievable at build time; US-only firm-size panel shown.",
}

us = {
    "question": "In the last two weeks, did this business use Artificial Intelligence (AI) in producing goods or services?",
    "survey": "Business Trends and Outlook Survey (BTOS)",
    "source_file": "AI Core Questions (legacy wording series)",
    "source_file_url": CENSUS_FILE_URL,
    "wording_change_pdf": CENSUS_WORDING_PDF,
    "pull_date": PULL_DATE,
    "wording_change": "2025-11-17",
    "new_series_start": "2025-12-04",
    "sept_points": [
        {"label": "Sept 2023", "fortnight": sept["2023_f"], "pct": sept["2023"]},
        {"label": "Sept 2024", "fortnight": sept["2024_f"], "pct": sept["2024"]},
        {"label": "Sept 2025", "fortnight": sept["2025_f"], "pct": sept["2025"]},
    ],
    "fortnightly": series,
    "sectors_sept_2025": sectors,
    "firm_size_sept_2025": firmsize,
}

meta = {
    "project": "canada-ai-adoption-inversion",
    "question": "On the one AI-adoption survey question Canada and the US asked businesses in nearly identical words, which country is actually ahead?",
    "headline_2025": {"canada_pct": 12.2, "canada_label": "Q2 2025", "us_pct": 10.0, "us_label": "Sept 2025"},
    "key_caveat": "Reference windows differ: StatCan asks about the 12 months preceding the survey; BTOS asks about the last two weeks. A longer window mechanically raises the share answering yes.",
    "pull_date": PULL_DATE,
    "license": "MIT",
    "sources": {
        "statcan_q2_2025": STATCAN_Q2_2025,
        "statcan_q2_2026": STATCAN_Q2_2026,
        "census_btos_downloads": CENSUS_FILE_URL,
        "census_wording_change": CENSUS_WORDING_PDF,
    },
}

(OUT / "meta.json").write_text(json.dumps(meta, indent=2))
(OUT / "canada.json").write_text(json.dumps(canada, indent=2))
(OUT / "us_legacy_series.json").write_text(json.dumps(us, indent=2))

# CSVs
with open(OUT / "us_fortnightly.csv", "w", newline="") as f:
    w = csv.writer(f)
    w.writerow(["fortnight", "collection_start", "reference_start", "reference_end", "pct_yes"])
    for pt in series:
        w.writerow([pt["fortnight"], pt["col_start"], pt["ref_start"], pt["ref_end"], pt["pct"]])
with open(OUT / "us_sectors_sept2025.csv", "w", newline="") as f:
    w = csv.writer(f)
    w.writerow(["naics", "sector", "pct_yes"])
    for s in sectors:
        w.writerow([s["naics"], s["sector"], s["pct"] if s["pct"] is not None else "S"])
with open(OUT / "us_firmsize_sept2025.csv", "w", newline="") as f:
    w = csv.writer(f)
    w.writerow(["employees_band", "pct_yes"])
    for s in firmsize:
        w.writerow([s["band"], s["pct"]])
with open(OUT / "canada_headline.csv", "w", newline="") as f:
    w = csv.writer(f)
    w.writerow(["quarter", "collection", "pct_yes", "source"])
    for h in canada["headline"]:
        w.writerow([h["quarter"], h["collection"], h["pct"], h.get("source", STATCAN_Q2_2025)])
with open(OUT / "canada_sectors_q2_2025.csv", "w", newline="") as f:
    w = csv.writer(f)
    w.writerow(["sector", "pct_2025", "pct_2024"])
    for s in canada["sectors_q2_2025"]:
        w.writerow([s["sector"], s["pct_2025"], s["pct_2024"] if s["pct_2024"] is not None else ""])

print("wrote", len(list(OUT.glob('*'))), "files to", OUT)
